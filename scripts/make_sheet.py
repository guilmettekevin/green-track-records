"""Create the Google Sheet starter workbook (Green Track Records.xlsx) from data/*.csv.

Upload the .xlsx to Google Drive and open it with Google Sheets
(or File > Import > Upload > "Replace spreadsheet").
"""
import csv
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(__file__), "..")
DATA = os.path.join(ROOT, "data")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "Green Track Records.xlsx")

FONT = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1A1A1A")
EDIT_FILL = PatternFill("solid", fgColor="FFF4E5")  # columns coaches are expected to fill in
TITLE = Font(name="Arial", size=16, bold=True, color="E35205")
BOLD = Font(name="Arial", size=10, bold=True)
SPARE = {"Results": 3000, "Athletes": 400, "Meets": 200, "Events": 20, "Coverage": 100, "Team Results": 60}  # blank rows pre-formatted as plain text so new marks are not turned into times/dates

TEXT_COLS = {
    "Results": ["Athlete", "Event", "Meet", "Mark", "Relay Total", "Status", "Notes", "Source", "Place", "Points", "Wind", "Series"],
    "Athletes": ["Athlete", "Merge Into", "Hide", "Notes"],
    "Meets": ["Meet", "Display Name", "Meet Type", "Notes"],
    "Events": ["Event", "Standard Event", "Display Name", "Leg", "Mark Min", "Mark Max", "Total Min", "Total Max"],
    "Coverage": ["Team", "Level", "Status", "Source", "Notes"],
    "Team Results": ["Team", "Level", "Meet", "Place", "Tied", "Points", "Teams", "Champion", "Champion Points", "Notes"],
}
WIDTHS = {"Gender": 9, "Athlete": 26, "Event": 22, "Meet": 38, "Date": 12, "Mark": 10, "Relay Total": 11,
          "Status": 9, "Notes": 30, "Source": 14, "Grad Year": 10, "Merge Into": 26, "Hide": 7,
          "Display Name": 36, "Meet Type": 13, "Standard Event": 22, "Type": 11, "Leg": 7, "Round": 8,
          "Hand Timed": 11, "Place": 7, "Points": 8, "Wind": 7, "Series": 40, "Mark Min": 10, "Mark Max": 10, "Total Min": 10, "Total Max": 10, "Order": 7,
          "Season": 8, "Team": 8, "Level": 10, "Green Results": 13, "Tied": 6, "Teams": 7, "Champion": 26, "Champion Points": 15}
EDITABLE = {"Results": {"Status", "Notes", "Place", "Points", "Wind", "Series"}, "Athletes": {"Grad Year", "Merge Into", "Hide", "Notes"},
            "Meets": {"Display Name", "Meet Type", "Notes"}, "Events": set(), "Coverage": {"Status", "Notes"},
            "Team Results": {"Season", "Team", "Level", "Meet", "Place", "Tied", "Points", "Teams", "Champion", "Champion Points", "Notes"}}

README = [
    ("Green High School Track & Field — Records Database", TITLE),
    ("This Google Sheet is the master copy. The website rebuilds itself from it every night "
     "(or right away with Track Records > Publish website now).", None),
    ("", None),
    ("TABS", BOLD),
    ("Results — one row per athlete per performance. Relays: one row per runner with the same "
     "Event, Meet, Date and Relay Total; the website groups them into a team. Runners not known yet: "
     "one row with Athlete = Unknown keeps the team on the lists, and the Data Check page lists it under "
     "Relays missing runners until all the runners are added.", None),
    ("Athletes — one row per athlete. Fill in Grad Year when you know it (unlocks class/grade stats). "
     "Merge Into: type the correct name to combine two spellings. Hide: Y to show 'Name withheld' on the site.", None),
    ("Meets — one row per meet name as entered. Display Name lets you tidy names "
     "(e.g. 'Lake' and 'Lake Meet' on the same date) without retyping results. Meet Type: Dual/Tri, Invitational, League, District, Regional, State, Indoor, Indoor State. District, Regional, State and Indoor State drive the postseason badges on athlete pages.", None),
    ("Coverage — which district, regional and state results are on file, one row per season, team and level. "
     "Filled in when official results are imported. Status: Complete, No Green athletes, Partial or Missing. "
     "Add a row (or change Status / Notes) to mark something you know is incomplete; the website's Data Check page "
     "turns this into a checklist of what's still to find.", None),
    ("Team Results — Green's team finishes the results files don't cover: older state meets, league or invitational "
     "titles. One row per season, team and level (District, Regional, State, League, or a meet name). Place 1 = champions; "
     "Tied: Y for a shared place. District, regional and state finishes from 2001 on are filled in automatically from the "
     "results files; a row here for the same season, team and level replaces the automatic one.", None),
    ("Events — maps each event name to a standard event. Mark Min/Max (and Total Min/Max for relays) "
     "are sanity limits: a mark better than the limit is held off the top lists until checked. Rarely needs editing.", None),
    ("", None),
    ("ADDING RESULTS", BOLD),
    ("Easiest: export 'All Career Performances' from Track Central, then Track Records > Import Track Central export. "
     "Only new rows are added, so you can import the whole file every time.", None),
    ("By hand: add a row at the bottom of Results. Name as Last, First. Date as 2026-04-18. "
     "Times like 11.24, 4:35.2, 10:05.33. Field marks like 18-4.5 (feet-inches). Mark 0-0 or blank = no mark.", None),
    ("Hand times: sprints and hurdles typed to tenths count as hand times on their own. For any other hand time "
     "(a dual-meet 1600 or 4x800, say), add h to the end: 4:38h. Last name only? Type it as given; the Data Check "
     "page lists those under Missing first names so they can be merged later.", None),
    ("Example row:  Boys | Doe, John | 1600M | Federal League Meet | 2026-05-08 | 4:35.21 | | | | Hand entry", None),
    ("Relay example (one row per runner):  Girls | Doe, Jane | 4X400M | District Meet | 2026-05-20 | 61.4 | 4:10.55", None),
    ("", None),
    ("FIXING THINGS", BOLD),
    ("Points: points the athlete scored for the team, as Track Central reported them. Filled in by importing a Track Central "
     "'Scoring Performances' report for a season. The website re-splits relay points evenly (1/4 each for a 4x relay) so every year counts the same.", None),
    ("Wind: the wind reading in m/s for the 100, 200, hurdles and long jump when the results list it (e.g. 1.4 or -0.3). "
     "Over +2.0 is wind-aided; the website shows it as 'w +2.4'.", None),
    ("Asterisk: end a Mark or Relay Total with * (e.g. 1:02.3h*) when the sources disagree; the site shows * with the row's Notes. "
     "Not sure which runner ran a relay leg? Add a row for each possibility with * in their Mark and explain in Notes.", None),
    ("Series: every attempt in a field event, when the score sheet has them. Throws and horizontal jumps: the marks in order, "
     "with F for a foul and P for a pass (e.g. 17-10, F, 16-7.25, 17-2.5). High jump and pole vault: each height with its "
     "attempts, O = make, X = miss, P = pass (e.g. 5-8 XO, 5-10 O, 6-0 XXX). The website shows the series when you tap the mark.", None),
    ("Series also holds lap splits for running events (800 and up) from fully automatic timing: each 400m lap in order, e.g. 1:22.43, 1:23.88, 1:25.96. "
     "They must add up to the Mark. The website shows them under a Splits button, with the running total.", None),
    ("Place: the finishing place in a final (1, 2, 3...). At the state meet, 1 = State champion and 1-8 = All-Ohio badges on the website. Leave it blank for prelims.", None),
    ("Status = Exclude hides a bad row from the site. Status = OK confirms a mark the site flagged. "
     "The website's Data Check page lists everything that looks off.", None),
    ("Shaded columns are the ones coaches normally edit. Please don't rename tabs or column headers — the website reads them by name.", None),
    ("Keep Mark and Relay Total formatted as Plain text (Format > Number > Plain text) so Sheets doesn't turn 4:35 into a clock time.", None),
    ("", None),
    (f"Starter data: Track Central 'All Career Performances' exports (boys and girls), printed 10/12/2026. Built {date.today():%m/%d/%Y}.", None),
]


def read(name):
    with open(os.path.join(DATA, name), newline="", encoding="utf-8") as f:
        r = list(csv.reader(f))
    return r[0], r[1:]


def add_tab(wb, title, header, data):
    ws = wb.create_sheet(title)
    ws.append(header)
    date_cols = {i for i, h in enumerate(header) if h == "Date"}
    num_cols = {i for i, h in enumerate(header) if h in ("Grad Year", "Order", "Season", "Green Results")}
    for row in data:
        out = []
        for i, v in enumerate(row):
            if i in date_cols and v:
                try:
                    v = date.fromisoformat(v)
                except ValueError:
                    pass
            elif i in num_cols and v.isdigit():
                v = int(v)
            out.append(v)
        ws.append(out)
    text_cols = set(TEXT_COLS.get(title, []))
    last = len(data) + 1 + SPARE.get(title, 0)
    for ci, h in enumerate(header, start=1):
        col = get_column_letter(ci)
        ws.column_dimensions[col].width = WIDTHS.get(h, 14)
        c = ws.cell(1, ci)
        c.font, c.fill = HEAD, HEAD_FILL
        fmt = "@" if h in text_cols else ("yyyy-mm-dd" if h == "Date" else None)
        fill = EDIT_FILL if h in EDITABLE.get(title, ()) else None
        for r in range(2, last + 1):
            cell = ws.cell(r, ci)
            cell.font = FONT
            if fmt:
                cell.number_format = fmt
            if fill:
                cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(data) + 1}"
    return ws, last


def main():
    wb = Workbook()
    rm = wb.active
    rm.title = "Read Me"
    rm.column_dimensions["A"].width = 130
    for i, (text, font) in enumerate(README, start=1):
        c = rm.cell(i, 1, text)
        c.font = font or FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")

    h, d = read("results.csv")
    ws, last = add_tab(wb, "Results", h, d)
    col = {name: get_column_letter(i + 1) for i, name in enumerate(h)}
    dv = DataValidation(type="list", formula1='"Boys,Girls"', allow_blank=True)
    dv.add(f"{col['Gender']}2:{col['Gender']}{last}")
    ws.add_data_validation(dv)
    dv = DataValidation(type="list", formula1='"Exclude,OK"', allow_blank=True)
    dv.add(f"{col['Status']}2:{col['Status']}{last}")
    ws.add_data_validation(dv)
    dv = DataValidation(type="list", formula1="=Events!$A$2:$A$200", allow_blank=True,
                        showErrorMessage=True, errorStyle="warning",
                        error="Not in the Events tab. Add it there first (or pick from the list).")
    dv.add(f"{col['Event']}2:{col['Event']}{last}")
    ws.add_data_validation(dv)

    h, d = read("athletes.csv")
    ws, last = add_tab(wb, "Athletes", h, d)
    dv = DataValidation(type="list", formula1='"Boys,Girls"', allow_blank=True)
    dv.add(f"A2:A{last}")
    ws.add_data_validation(dv)
    dv = DataValidation(type="list", formula1='"Y"', allow_blank=True)
    dv.add(f"E2:E{last}")
    ws.add_data_validation(dv)

    h, d = read("meets.csv")
    ws, last = add_tab(wb, "Meets", h, d)
    dv = DataValidation(type="list", formula1='"Dual/Tri,Invitational,League,District,Regional,State,Indoor,Indoor State"', allow_blank=True)
    dv.add(f"D2:D{last}")
    ws.add_data_validation(dv)

    h, d = read("events.csv")
    add_tab(wb, "Events", h, d)

    if os.path.exists(os.path.join(DATA, "coverage.csv")):
        h, d = read("coverage.csv")
        ws, last = add_tab(wb, "Coverage", h, d)
        col = {name: get_column_letter(i + 1) for i, name in enumerate(h)}
        for name, choices in (("Team", "Boys,Girls"), ("Level", "League,District,Regional,State,Points"),
                              ("Status", "Complete,No Green athletes,Partial,Missing")):
            dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
            dv.add(f"{col[name]}2:{col[name]}{last}")
            ws.add_data_validation(dv)

    if os.path.exists(os.path.join(DATA, "team.csv")):
        h, d = read("team.csv")
        ws, last = add_tab(wb, "Team Results", h, d)
        col = {name: get_column_letter(i + 1) for i, name in enumerate(h)}
        for name, choices in (("Team", "Boys,Girls"), ("Tied", "Y")):
            dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
            dv.add(f"{col[name]}2:{col[name]}{last}")
            ws.add_data_validation(dv)

    wb.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    main()
